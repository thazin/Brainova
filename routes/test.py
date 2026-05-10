from flask import Blueprint, render_template, request, redirect, url_for, session, flash, send_file, abort
from database import get_db
from functools import wraps
import io
import random
from datetime import datetime

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.enums import TA_CENTER, TA_LEFT

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

test_bp = Blueprint('test', __name__)


def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if 'user_id' not in session:
            flash('Please log in to take the test.', 'warning')
            return redirect(url_for('auth.login'))
        return f(*args, **kwargs)
    return decorated


def calculate_iq(score, total):
    """Map raw score percentage to an IQ score (70–145 range)."""
    if total == 0:
        return 100
    pct = score / total
    # IQ: 70 at 0%, 100 at 50%, 145 at 100%
    if pct <= 0.5:
        return int(70 + pct * 60)
    else:
        return int(100 + (pct - 0.5) * 90)


@test_bp.route('/')
@login_required
def start():
    db = get_db()
    counts = {}
    for cat in ['logical', 'verbal', 'numerical', 'spatial']:
        counts[cat] = db.execute(
            'SELECT COUNT(*) FROM questions WHERE category = ?', (cat,)
        ).fetchone()[0]
    total = db.execute('SELECT COUNT(*) FROM questions').fetchone()[0]
    last_result = db.execute(
        'SELECT * FROM test_results WHERE user_id = ? ORDER BY taken_at DESC LIMIT 1',
        (session['user_id'],)
    ).fetchone()
    db.close()
    return render_template('test/start.html', counts=counts, total=total, last_result=last_result)


@test_bp.route('/take', methods=['GET', 'POST'])
@login_required
def take():
    if request.method == 'POST':
        # Score the submitted answers
        db = get_db()
        questions = db.execute('SELECT id FROM questions').fetchall()
        q_ids = [q['id'] for q in questions]

        score = 0
        answers_log = []
        for q_id in q_ids:
            selected_id = request.form.get(f'q_{q_id}')
            is_correct = 0
            selected_int = None
            if selected_id:
                selected_int = int(selected_id)
                correct_ans = db.execute(
                    'SELECT id FROM answers WHERE question_id = ? AND is_correct = 1',
                    (q_id,)
                ).fetchone()
                if correct_ans and correct_ans['id'] == selected_int:
                    is_correct = 1
                    score += 1
            answers_log.append((q_id, selected_int, is_correct))

        total = len(q_ids)
        iq = calculate_iq(score, total)

        cur = db.execute(
            'INSERT INTO test_results (user_id, score, total_questions, iq_score) VALUES (?, ?, ?, ?)',
            (session['user_id'], score, total, iq)
        )
        result_id = cur.lastrowid
        for q_id, sel_id, is_correct in answers_log:
            db.execute(
                'INSERT INTO test_answers (result_id, question_id, selected_answer_id, is_correct) VALUES (?, ?, ?, ?)',
                (result_id, q_id, sel_id, is_correct)
            )
        db.commit()
        db.close()
        return redirect(url_for('test.result', result_id=result_id))

    # GET — load all questions with their options
    db = get_db()
    questions = db.execute('SELECT * FROM questions ORDER BY RANDOM()').fetchall()
    q_list = []
    for q in questions:
        opts = db.execute(
            'SELECT * FROM answers WHERE question_id = ?', (q['id'],)
        ).fetchall()
        q_list.append({'question': q, 'options': opts})
    db.close()

    if not q_list:
        flash('No questions available. Please ask an admin to add questions.', 'warning')
        return redirect(url_for('test.start'))

    return render_template('test/take.html', q_list=q_list)


@test_bp.route('/result/<int:result_id>')
@login_required
def result(result_id):
    db = get_db()
    res = db.execute(
        'SELECT * FROM test_results WHERE id = ? AND user_id = ?',
        (result_id, session['user_id'])
    ).fetchone()
    if not res:
        db.close()
        flash('Result not found.', 'danger')
        return redirect(url_for('test.start'))

    # Load each question with selected answer and correct answer
    details = db.execute(
        '''SELECT ta.is_correct, ta.selected_answer_id,
                  q.question_text, q.category,
                  sa.option_text AS selected_text,
                  ca.option_text AS correct_text
           FROM test_answers ta
           JOIN questions q ON ta.question_id = q.id
           LEFT JOIN answers sa ON ta.selected_answer_id = sa.id
           JOIN answers ca ON ca.question_id = q.id AND ca.is_correct = 1
           WHERE ta.result_id = ?
           ORDER BY ta.id''',
        (result_id,)
    ).fetchall()
    db.close()

    iq_label = _iq_label(res['iq_score'])
    return render_template('test/result.html', res=res, details=details, iq_label=iq_label)


def _iq_label(iq):
    if iq >= 130:
        return ('Very Superior', 'success')
    elif iq >= 120:
        return ('Superior', 'success')
    elif iq >= 110:
        return ('High Average', 'info')
    elif iq >= 90:
        return ('Average', 'primary')
    elif iq >= 80:
        return ('Low Average', 'warning')
    else:
        return ('Below Average', 'danger')


def _load_result_for_user(result_id, user_id):
    db = get_db()
    res = db.execute(
        '''SELECT tr.*, u.username, u.email
           FROM test_results tr
           JOIN users u ON tr.user_id = u.id
           WHERE tr.id = ? AND tr.user_id = ?''',
        (result_id, user_id)
    ).fetchone()
    if not res:
        db.close()
        return None, None
    details = db.execute(
        '''SELECT ta.is_correct, ta.selected_answer_id,
                  q.question_text, q.category,
                  sa.option_text AS selected_text,
                  ca.option_text AS correct_text
           FROM test_answers ta
           JOIN questions q ON ta.question_id = q.id
           LEFT JOIN answers sa ON ta.selected_answer_id = sa.id
           JOIN answers ca ON ca.question_id = q.id AND ca.is_correct = 1
           WHERE ta.result_id = ?
           ORDER BY ta.id''',
        (result_id,)
    ).fetchall()
    db.close()
    return res, details


@test_bp.route('/result/<int:result_id>/delete', methods=['POST'])
@login_required
def delete_result(result_id):
    db = get_db()
    res = db.execute(
        'SELECT id FROM test_results WHERE id = ? AND user_id = ?',
        (result_id, session['user_id'])
    ).fetchone()
    if not res:
        db.close()
        flash('Result not found.', 'danger')
        return redirect(url_for('profile.view'))

    db.execute('DELETE FROM test_answers WHERE result_id = ?', (result_id,))
    db.execute(
        'DELETE FROM test_results WHERE id = ? AND user_id = ?',
        (result_id, session['user_id'])
    )
    db.commit()
    db.close()
    flash('Test result deleted.', 'success')
    return redirect(url_for('profile.view'))


@test_bp.route('/result/<int:result_id>/download/pdf')
@login_required
def download_pdf(result_id):
    res, details = _load_result_for_user(result_id, session['user_id'])
    if not res:
        flash('Result not found.', 'danger')
        return redirect(url_for('test.start'))

    label, _ = _iq_label(res['iq_score'])
    accuracy = round((res['score'] / res['total_questions']) * 100, 1) if res['total_questions'] else 0
    taken_at = res['taken_at']

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=2*cm, rightMargin=2*cm,
        topMargin=1.8*cm, bottomMargin=1.8*cm,
        title=f'Brainova Result #{result_id}'
    )
    styles = getSampleStyleSheet()
    h_title = ParagraphStyle('Title', parent=styles['Title'], fontSize=22, alignment=TA_CENTER, spaceAfter=6)
    h_sub = ParagraphStyle('Sub', parent=styles['Normal'], fontSize=11, alignment=TA_CENTER, textColor=colors.grey, spaceAfter=18)
    h_section = ParagraphStyle('Section', parent=styles['Heading2'], fontSize=14, spaceBefore=12, spaceAfter=8)
    body = styles['BodyText']

    story = []
    story.append(Paragraph('Brainova Result Report', h_title))
    story.append(Paragraph(f"Generated for {res['username']} ({res['email']})", h_sub))
    story.append(HRFlowable(width='100%', thickness=1, color=colors.lightgrey, spaceAfter=12))

    summary_data = [
        ['IQ Score', str(res['iq_score'])],
        ['Classification', label],
        ['Correct Answers', f"{res['score']} / {res['total_questions']}"],
        ['Accuracy', f"{accuracy}%"],
        ['Taken At', str(taken_at)],
    ]
    summary_tbl = Table(summary_data, colWidths=[5*cm, 10*cm])
    summary_tbl.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#f1f3f5')),
        ('TEXTCOLOR', (0, 0), (0, -1), colors.HexColor('#333333')),
        ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 11),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ('TOPPADDING', (0, 0), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.lightgrey),
    ]))
    story.append(summary_tbl)

    story.append(Paragraph('Answer Review', h_section))

    review_data = [['#', 'Category', 'Question', 'Your Answer', 'Correct Answer', 'Result']]
    for idx, d in enumerate(details, 1):
        review_data.append([
            str(idx),
            d['category'].capitalize(),
            Paragraph(d['question_text'], body),
            Paragraph(d['selected_text'] or '—', body),
            Paragraph(d['correct_text'], body),
            'Correct' if d['is_correct'] else 'Wrong',
        ])

    review_tbl = Table(
        review_data,
        colWidths=[0.8*cm, 2.2*cm, 5.5*cm, 3.2*cm, 3.2*cm, 1.7*cm],
        repeatRows=1
    )
    style = TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0d6efd')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.lightgrey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#fafafa')]),
    ])
    for i, d in enumerate(details, 1):
        color = colors.HexColor('#198754') if d['is_correct'] else colors.HexColor('#dc3545')
        style.add('TEXTCOLOR', (5, i), (5, i), color)
        style.add('FONTNAME', (5, i), (5, i), 'Helvetica-Bold')
    review_tbl.setStyle(style)
    story.append(review_tbl)

    story.append(Spacer(1, 18))
    story.append(Paragraph(
        f'<para align="center"><font size="8" color="#888888">'
        f'Report generated by Brainova · {datetime.now().strftime("%Y-%m-%d %H:%M")}'
        f'</font></para>', body
    ))

    doc.build(story)
    buf.seek(0)
    filename = f"iq_result_{result_id}_{res['username']}.pdf"
    return send_file(buf, mimetype='application/pdf', as_attachment=True, download_name=filename)


@test_bp.route('/result/<int:result_id>/download/excel')
@login_required
def download_excel(result_id):
    res, details = _load_result_for_user(result_id, session['user_id'])
    if not res:
        flash('Result not found.', 'danger')
        return redirect(url_for('test.start'))

    label, _ = _iq_label(res['iq_score'])
    accuracy = round((res['score'] / res['total_questions']) * 100, 1) if res['total_questions'] else 0

    wb = openpyxl.Workbook()
    thin = Side(border_style='thin', color='DDDDDD')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill('solid', fgColor='0D6EFD')
    header_font = Font(bold=True, color='FFFFFF', size=11)
    label_fill = PatternFill('solid', fgColor='F1F3F5')
    label_font = Font(bold=True, color='333333')

    # Sheet 1 — Summary
    ws = wb.active
    ws.title = 'Summary'
    ws['A1'] = 'Brainova Result Report'
    ws['A1'].font = Font(bold=True, size=16, color='0D6EFD')
    ws.merge_cells('A1:B1')
    ws['A1'].alignment = Alignment(horizontal='center')

    rows = [
        ('User', res['username']),
        ('Email', res['email']),
        ('IQ Score', res['iq_score']),
        ('Classification', label),
        ('Correct Answers', f"{res['score']} / {res['total_questions']}"),
        ('Accuracy (%)', accuracy),
        ('Taken At', str(res['taken_at'])),
    ]
    for i, (k, v) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=k).font = label_font
        ws.cell(row=i, column=1).fill = label_fill
        ws.cell(row=i, column=1).border = border
        ws.cell(row=i, column=2, value=v).border = border
    ws.column_dimensions['A'].width = 22
    ws.column_dimensions['B'].width = 38

    # Sheet 2 — Answer Review
    ws2 = wb.create_sheet('Answer Review')
    headers = ['#', 'Category', 'Question', 'Your Answer', 'Correct Answer', 'Result']
    for col, h in enumerate(headers, 1):
        c = ws2.cell(row=1, column=col, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = border

    correct_fill = PatternFill('solid', fgColor='D1E7DD')
    wrong_fill = PatternFill('solid', fgColor='F8D7DA')

    for idx, d in enumerate(details, 1):
        row = idx + 1
        values = [
            idx,
            d['category'].capitalize(),
            d['question_text'],
            d['selected_text'] or '—',
            d['correct_text'],
            'Correct' if d['is_correct'] else 'Wrong',
        ]
        for col, v in enumerate(values, 1):
            c = ws2.cell(row=row, column=col, value=v)
            c.border = border
            c.alignment = Alignment(wrap_text=True, vertical='top')
        result_cell = ws2.cell(row=row, column=6)
        result_cell.fill = correct_fill if d['is_correct'] else wrong_fill
        result_cell.font = Font(bold=True, color='0F5132' if d['is_correct'] else '842029')
        result_cell.alignment = Alignment(horizontal='center', vertical='center')

    widths = [5, 14, 55, 28, 28, 12]
    for i, w in enumerate(widths, 1):
        ws2.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws2.freeze_panes = 'A2'

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"iq_result_{result_id}_{res['username']}.xlsx"
    return send_file(
        buf,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename
    )


# ---------------------------------------------------------------------------
# Admin-defined published tests
# ---------------------------------------------------------------------------

def _select_questions_for_test(db, test_row):
    """Pick `num_questions` from the configured pool, optionally shuffled.

    Done in Python (not SQL ORDER BY RANDOM()) so the pool filter and shuffle
    flag are always respected the same way across DB backends.
    """
    if test_row['category']:
        rows = db.execute(
            'SELECT * FROM questions WHERE category = ?',
            (test_row['category'],)
        ).fetchall()
    else:
        rows = db.execute('SELECT * FROM questions').fetchall()

    pool = list(rows)
    if test_row['shuffle']:
        random.shuffle(pool)
    return pool[: test_row['num_questions']]


@test_bp.route('/list')
@login_required
def list_published():
    """Catalog of admin-published tests visible to a normal user."""
    db = get_db()
    rows = db.execute(
        '''SELECT t.*, u.username AS created_by_name
           FROM tests t
           JOIN users u ON t.created_by = u.id
           WHERE t.status = 'published'
           ORDER BY COALESCE(t.published_at, t.created_at) DESC'''
    ).fetchall()
    db.close()
    return render_template('test/list.html', tests=rows)


@test_bp.route('/<int:test_id>/begin', methods=['GET', 'POST'])
@login_required
def begin(test_id):
    db = get_db()
    test = db.execute(
        "SELECT * FROM tests WHERE id = ? AND status = 'published'",
        (test_id,)
    ).fetchone()
    if not test:
        db.close()
        # Hide draft/unpublished tests from users entirely.
        abort(404)

    if request.method == 'POST':
        # The form posts back the exact set of question ids it rendered, so
        # scoring uses the same questions the user actually saw (important
        # when shuffle=True).
        q_ids = [int(x) for x in request.form.getlist('q_ids') if x.isdigit()]

        score = 0
        answers_log = []
        for q_id in q_ids:
            selected_id = request.form.get(f'q_{q_id}')
            is_correct = 0
            selected_int = None
            if selected_id:
                selected_int = int(selected_id)
                correct_ans = db.execute(
                    'SELECT id FROM answers WHERE question_id = ? AND is_correct = 1',
                    (q_id,)
                ).fetchone()
                if correct_ans and correct_ans['id'] == selected_int:
                    is_correct = 1
                    score += 1
            answers_log.append((q_id, selected_int, is_correct))

        total = len(q_ids)
        iq = calculate_iq(score, total)
        cur = db.execute(
            '''INSERT INTO test_results (user_id, test_id, score,
                                         total_questions, iq_score)
               VALUES (?, ?, ?, ?, ?)''',
            (session['user_id'], test_id, score, total, iq)
        )
        result_id = cur.lastrowid
        for q_id, sel_id, is_correct in answers_log:
            db.execute(
                '''INSERT INTO test_answers (result_id, question_id,
                                             selected_answer_id, is_correct)
                   VALUES (?, ?, ?, ?)''',
                (result_id, q_id, sel_id, is_correct)
            )
        db.commit()
        db.close()
        return redirect(url_for('test.result', result_id=result_id))

    questions = _select_questions_for_test(db, test)
    if not questions:
        db.close()
        flash('This test has no questions available right now.', 'warning')
        return redirect(url_for('test.list_published'))

    q_list = []
    for q in questions:
        opts = db.execute(
            'SELECT * FROM answers WHERE question_id = ?', (q['id'],)
        ).fetchall()
        # Shuffle answer options per-question when the test is randomized so
        # the correct answer doesn't always sit in the same slot.
        opts = list(opts)
        if test['shuffle']:
            random.shuffle(opts)
        q_list.append({'question': q, 'options': opts})
    db.close()

    return render_template('test/take_custom.html', test=test, q_list=q_list)


@test_bp.route('/<int:test_id>/leaderboard')
@login_required
def leaderboard(test_id):
    """Monthly top scorers for a single published test.

    Only counts attempts from the current calendar month — the ranking
    automatically resets at the start of each month. Past results remain in
    the database (for the user's own history & exports) but stop appearing
    on the leaderboard once the month rolls over.
    """
    db = get_db()
    test = db.execute(
        "SELECT * FROM tests WHERE id = ? AND status = 'published'",
        (test_id,)
    ).fetchone()
    if not test:
        db.close()
        abort(404)

    # Window: first day of the current month (inclusive) → first day of next
    # month (exclusive). SQLite stores taken_at as 'YYYY-MM-DD HH:MM:SS', so
    # an ISO-style string comparison works correctly.
    now = datetime.now()
    month_start = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    if month_start.month == 12:
        next_month = month_start.replace(year=month_start.year + 1, month=1)
    else:
        next_month = month_start.replace(month=month_start.month + 1)
    month_start_s = month_start.strftime('%Y-%m-%d %H:%M:%S')
    next_month_s = next_month.strftime('%Y-%m-%d %H:%M:%S')
    period_label = month_start.strftime('%B %Y')

    # Best result per user *within this month*, ranked by IQ then earliest
    # taken_at to break ties. The correlated `attempts` subquery counts every
    # attempt the user made on this test in-month, not just the best one.
    rows = db.execute(
        '''SELECT tr.id AS result_id,
                  tr.user_id,
                  u.username,
                  tr.score,
                  tr.total_questions,
                  tr.iq_score,
                  tr.taken_at,
                  (SELECT COUNT(*) FROM test_results x
                     WHERE x.user_id = tr.user_id
                       AND x.test_id = ?
                       AND x.taken_at >= ? AND x.taken_at < ?
                  ) AS attempts
           FROM test_results tr
           JOIN users u ON u.id = tr.user_id
           JOIN (
                SELECT user_id, MAX(iq_score) AS best_iq
                FROM test_results
                WHERE test_id = ?
                  AND taken_at >= ? AND taken_at < ?
                GROUP BY user_id
           ) best
             ON best.user_id = tr.user_id
            AND best.best_iq = tr.iq_score
           WHERE tr.test_id = ?
             AND tr.taken_at >= ? AND tr.taken_at < ?
           GROUP BY tr.user_id
           ORDER BY tr.iq_score DESC, MIN(tr.taken_at) ASC''',
        (test_id, month_start_s, next_month_s,
         test_id, month_start_s, next_month_s,
         test_id, month_start_s, next_month_s)
    ).fetchall()
    db.close()

    def _ordinal(n):
        # 11/12/13 are 'th' regardless of last digit; 1/2/3 → st/nd/rd; rest → th.
        if 10 <= n % 100 <= 20:
            suffix = 'th'
        else:
            suffix = {1: 'st', 2: 'nd', 3: 'rd'}.get(n % 10, 'th')
        return f"{n}{suffix}"

    # Competition ranking ("1224"): users with the same IQ share the same
    # rank; the next distinct IQ skips ahead by the size of the tied group.
    # The SQL ORDER BY still uses earliest taken_at as a stable secondary
    # sort so the *display order* within a tied group is deterministic, but
    # the rank number itself reflects the tie.
    me_id = session.get('user_id')
    entries = []
    my_rank = None
    prev_iq = None
    current_rank = 0
    for idx, r in enumerate(rows, 1):
        if r['iq_score'] != prev_iq:
            current_rank = idx
            prev_iq = r['iq_score']

        accuracy = round((r['score'] / r['total_questions']) * 100, 1) if r['total_questions'] else 0
        entries.append({
            'rank': current_rank,
            'rank_label': _ordinal(current_rank),
            'result_id': r['result_id'],
            'user_id': r['user_id'],
            'username': r['username'],
            'score': r['score'],
            'total_questions': r['total_questions'],
            'iq_score': r['iq_score'],
            'taken_at': r['taken_at'],
            'attempts': r['attempts'],
            'accuracy': accuracy,
            'is_me': r['user_id'] == me_id,
        })
        if r['user_id'] == me_id:
            my_rank = current_rank

    # Flag ties so the template can mark them visually (e.g. "T-1").
    rank_counts = {}
    for e in entries:
        rank_counts[e['rank']] = rank_counts.get(e['rank'], 0) + 1
    for e in entries:
        e['tied'] = rank_counts[e['rank']] > 1
        e['tied_with'] = rank_counts[e['rank']] - 1  # how many others share this rank

    # Pick one representative per podium tier (rank 1, 2, 3). If multiple
    # users tie at rank 1, the others show up in the table beneath; the
    # podium just notes "+N tied" so the layout stays clean.
    def _pick(rank_value):
        for e in entries:
            if e['rank'] == rank_value:
                return e
        return None

    podium = {
        'first':  _pick(1),
        'second': _pick(2),
        'third':  _pick(3),
    }

    return render_template(
        'test/leaderboard.html',
        test=test,
        entries=entries,
        my_rank=my_rank,
        total_players=len(entries),
        period_label=period_label,
        podium=podium,
    )
